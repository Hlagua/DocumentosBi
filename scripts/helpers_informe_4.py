import os
import sys
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from latex2mathml.converter import convert as latex_to_mathml
import lxml.etree as ET

# ---------------------------------------------------------------------------
# Setup LaTeX to OMML converter
# ---------------------------------------------------------------------------
XSL_PATH = r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'
if os.path.exists(XSL_PATH):
    xslt_doc = ET.parse(XSL_PATH)
    xslt_transform = ET.XSLT(xslt_doc)
else:
    xslt_transform = None

def get_omml(latex_code):
    if not xslt_transform:
        return None
    try:
        cleaned = latex_code.strip()
        cleaned = re.sub(r'\\tag\{.*?\}', '', cleaned)
        cleaned = cleaned.replace('\u00a0', ' ')
        mml = latex_to_mathml(cleaned)
        tree = ET.fromstring(mml.encode('utf-8'))
        omml_tree = xslt_transform(tree)
        return ET.tostring(omml_tree, encoding='utf-8').decode('utf-8')
    except Exception:
        return None

# ---------------------------------------------------------------------------
# Document Initialization with Institutional Layout
# ---------------------------------------------------------------------------
def init_document():
    doc = docx.Document()
    
    section = doc.sections[0]
    section.page_width = Inches(8.27)   # A4 Width
    section.page_height = Inches(11.69) # A4 Height
    section.top_margin = Inches(0.98)   # 2.5 cm
    section.bottom_margin = Inches(0.98)# 2.5 cm
    section.left_margin = Inches(1.18)  # 3.0 cm
    section.right_margin = Inches(0.98) # 2.5 cm
    section.header_distance = Inches(0.49)
    section.footer_distance = Inches(0.49)
    section.different_first_page_header_footer = False
    
    # Header: 3-column table
    header = section.header
    hp = header.paragraphs[0]
    hp.text = ""
    hp.paragraph_format.space_before = Pt(0)
    hp.paragraph_format.space_after = Pt(0)
    
    htbl = header.add_table(1, 3, Inches(6.11))
    htbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    htbl.autofit = False
    
    col_widths = [Inches(0.85), Inches(4.41), Inches(0.85)]
    for i, w in enumerate(col_widths):
        htbl.rows[0].cells[i].width = w
        htbl.rows[0].cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        
    # Left cell: Logo UTA
    p_left = htbl.rows[0].cells[0].paragraphs[0]
    p_left.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_left.paragraph_format.space_before = Pt(0)
    p_left.paragraph_format.space_after = Pt(0)
    if os.path.exists('img/logo_uta.jpg'):
        p_left.add_run().add_picture('img/logo_uta.jpg', width=Inches(0.75))
        
    # Center cell: Institutional text
    p_center = htbl.rows[0].cells[1].paragraphs[0]
    p_center.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_center.paragraph_format.line_spacing = 1.05
    p_center.paragraph_format.space_before = Pt(0)
    p_center.paragraph_format.space_after = Pt(0)
    
    lines = [
        ("UNIVERSIDAD TÉCNICA DE AMBATO", 10.5, True),
        ("FACULTAD DE INGENIERÍA EN SISTEMAS ELECTRÓNICA E INDUSTRIAL", 8.0, True),
        ("CARRERA DE SOFTWARE", 8.0, True),
        ("CICLO ACADÉMICO: JULIO – DICIEMBRE 2026", 8.0, True),
    ]
    for idx, (t, sz, b) in enumerate(lines):
        if idx > 0:
            r = p_center.add_run("\n" + t)
        else:
            r = p_center.add_run(t)
        r.font.name = "Times New Roman"
        r.font.size = Pt(sz)
        r.bold = b
        
    # Right cell: Logo FISEI
    p_right = htbl.rows[0].cells[2].paragraphs[0]
    p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_right.paragraph_format.space_before = Pt(0)
    p_right.paragraph_format.space_after = Pt(0)
    if os.path.exists('img/logo_fisei.jpg'):
        p_right.add_run().add_picture('img/logo_fisei.jpg', width=Inches(0.75))
        
    # Divider rule under header
    p_rule = header.add_paragraph()
    p_rule.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_rule.paragraph_format.space_before = Pt(2)
    p_rule.paragraph_format.space_after = Pt(0)
    pBrd = parse_xml(r'<w:pBrd %s><w:bottom w:val="single" w:sz="6" w:space="1" w:color="000000"/></w:pBrd>' % nsdecls('w'))
    p_rule._p.get_or_add_pPr().append(pBrd)
    
    # Footer: Centered page number
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    frun = fp.add_run()
    frun.font.name = "Times New Roman"
    frun.font.size = Pt(8.5)
    fldSimple = parse_xml(r'<w:fldSimple %s w:instr="PAGE"/>' % nsdecls('w'))
    fp._p.append(fldSimple)
    
    return doc

# ---------------------------------------------------------------------------
# Formatting Helpers (Clean, Simple, Accessible)
# ---------------------------------------------------------------------------
def clean_math_in_prose(text):
    text = re.sub(r'\\text\{([^}]+)\}', r'\1', text)
    text = re.sub(r'\\mathbf\{([^}]+)\}', r'\1', text)
    text = re.sub(r'\\hat\{([^}]+)\}', r'\1^', text)
    text = re.sub(r'\\overline\{([^}]+)\}', r'\1', text)
    text = re.sub(r'\\bar\{([^}]+)\}', r'\1', text)
    text = re.sub(r'\\Delta\s*', 'Δ', text)
    text = re.sub(r'\\pm\s*', '±', text)
    text = re.sub(r'\\le\s*', '≤', text)
    text = re.sub(r'\\ge\s*', '≥', text)
    text = re.sub(r'\\times\s*', '×', text)
    text = re.sub(r'\\cdot\s*', '·', text)
    text = re.sub(r'\\approx\s*', '≈', text)
    text = re.sub(r'\\in\s*', '∈', text)
    text = re.sub(r'\\chi\^2\s*', 'χ²', text)
    text = re.sub(r'\\sigma\^2\s*', 'σ²', text)
    text = re.sub(r'\\sigma\s*', 'σ', text)
    text = re.sub(r'\\rho\s*', 'ρ', text)
    text = re.sub(r'\\cos\s*', 'cos', text)
    text = re.sub(r'\\ln\s*', 'ln', text)
    text = text.replace('$', '')
    return text

def render_inline(p, text, base_size=11.0, default_italic=False):
    text = clean_math_in_prose(text)
    text = re.sub(r'<br\s*/?>', '\n', text)
    
    pattern = re.compile(
        r'(?P<bold>\*\*(.*?)\*\*)|'
        r'(?P<italic>(?<!\*)\*(?!\*)(.*?)(?<!\*)\*(?!\*))|'
        r'(?P<code>`([^`]+)`)|'
        r'(?P<link>\[(.*?)\]\((.*?)\))'
    )
    
    def add_text_run(txt, bold=False, italic=default_italic, code=False):
        r = p.add_run(txt)
        r.font.name = "Times New Roman"
        r.font.size = Pt(base_size)
        r.bold = bold
        r.italic = italic
        r.font.color.rgb = RGBColor(0, 0, 0)
        
    last_idx = 0
    for match in pattern.finditer(text):
        start, end = match.span()
        if start > last_idx:
            chunk = text[last_idx:start]
            if chunk:
                add_text_run(chunk)
                
        kind = match.lastgroup
        if kind == 'bold':
            add_text_run(match.group(2), bold=True)
        elif kind == 'italic':
            add_text_run(match.group(4), italic=True)
        elif kind == 'code':
            add_text_run(match.group(6), italic=True)
        elif kind == 'link':
            add_text_run(match.group(8))
            
        last_idx = end
        
    if last_idx < len(text):
        chunk = text[last_idx:]
        if chunk:
            add_text_run(chunk)

def add_heading(doc, level, text):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    
    if level == 1:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        render_inline(p, text, base_size=14.0)
        p.runs[0].bold = True
    elif level == 2:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        render_inline(p, text, base_size=12.5)
        p.runs[0].bold = True
    elif level == 3:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(9)
        p.paragraph_format.space_after = Pt(3)
        render_inline(p, text, base_size=11.5)
        p.runs[0].bold = True
    elif level == 4:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(2)
        render_inline(p, text, base_size=11.0)
        p.runs[0].bold = True
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        render_inline(p, text, base_size=11.0)
        if p.runs:
            p.runs[0].bold = True

def add_paragraph(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4.5)
    render_inline(p, text, base_size=11.0)

def add_bullet(doc, text, bold_prefix=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3.5)
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.28)
    
    # Strip any accidental double numbering or leading bullet characters
    if bold_prefix:
        bold_prefix = re.sub(r'^[\s\*\-\•\u2022\d+\.\)]+', '', bold_prefix).strip()
    text = re.sub(r'^[\s\*\-\•\u2022\d+\.\)]+', '', text).strip()
    
    r_bullet = p.add_run("•  ")
    r_bullet.font.name = "Times New Roman"
    r_bullet.font.size = Pt(11.0)
    r_bullet.bold = True
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + " ")
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11.0)
        r_pre.bold = True
        
    render_inline(p, text, base_size=11.0)

def add_numbered(doc, num, text, bold_prefix=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3.5)
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.28)
    
    # Strip any accidental duplicate numbers or bullets
    if bold_prefix:
        bold_prefix = re.sub(r'^[\s\*\-\•\u2022\d+\.\)]+', '', bold_prefix).strip()
    text = re.sub(r'^[\s\*\-\•\u2022\d+\.\)]+', '', text).strip()
    
    r_num = p.add_run(f"{num}.  ")
    r_num.font.name = "Times New Roman"
    r_num.font.size = Pt(11.0)
    r_num.bold = True
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + " ")
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11.0)
        r_pre.bold = True
        
    render_inline(p, text, base_size=11.0)

def add_checkbox(doc, checked, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.28)
    
    box = "☒  " if checked else "☐  "
    r_box = p.add_run(box)
    r_box.font.name = "Segoe UI Symbol"
    r_box.font.size = Pt(11.0)
    r_box.bold = True
    
    render_inline(p, text, base_size=11.0)

def add_block_math(doc, latex_code):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    omml = get_omml(latex_code)
    if omml:
        try:
            if not omml.startswith('<m:oMathPara'):
                omml_wrapped = f'<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">{omml}</m:oMathPara>'
            else:
                omml_wrapped = omml
            p._p.append(parse_xml(omml_wrapped))
            return
        except Exception:
            try:
                p._p.append(parse_xml(omml))
                return
            except Exception:
                pass
    r = p.add_run(clean_math_in_prose(latex_code))
    r.font.name = "Times New Roman"
    r.font.size = Pt(10.5)
    r.italic = True

def add_figure(doc, img_path, title_text, width_inches=4.9):
    if not os.path.exists(img_path):
        print(f"WARNING: Image not found: {img_path}")
        return
        
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(10)
    p_img.paragraph_format.space_after = Pt(3)
    p_img.paragraph_format.keep_with_next = True
    
    r_img = p_img.add_run()
    r_img.add_picture(img_path, width=Inches(width_inches))
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(10)
    p_cap.paragraph_format.line_spacing = 1.05
    
    render_inline(p_cap, title_text, base_size=10.0)
    if p_cap.runs:
        m = re.match(r'^(Figura\s+\d+:?)(.*)$', p_cap.runs[0].text)
        if m:
            p_cap.runs[0].text = m.group(1)
            p_cap.runs[0].bold = True
            r_rest = p_cap.add_run(m.group(2))
            r_rest.font.name = "Times New Roman"
            r_rest.font.size = Pt(10.0)

def add_table(doc, title_text, headers, rows, col_widths=None):
    if title_text:
        p_t = doc.add_paragraph()
        p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_t.paragraph_format.space_before = Pt(10)
        p_t.paragraph_format.space_after = Pt(4)
        p_t.paragraph_format.keep_with_next = True
        
        render_inline(p_t, title_text, base_size=10.5)
        if p_t.runs:
            m = re.match(r'^(Tabla\s+\d+:?)(.*)$', p_t.runs[0].text)
            if m:
                p_t.runs[0].text = m.group(1)
                p_t.runs[0].bold = True
                r_rest = p_t.add_run(m.group(2))
                r_rest.font.name = "Times New Roman"
                r_rest.font.size = Pt(10.5)
            
    num_cols = len(headers)
    num_rows = len(rows) + 1
    
    tbl = doc.add_table(num_rows, num_cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    tblPr = tbl._tbl.tblPr
    tblCellMar = parse_xml(r'''<w:tblCellMar %s>
        <w:top w:w="80" w:type="dxa"/>
        <w:bottom w:w="80" w:type="dxa"/>
        <w:left w:w="110" w:type="dxa"/>
        <w:right w:w="110" w:type="dxa"/>
    </w:tblCellMar>''' % nsdecls('w'))
    tblPr.append(tblCellMar)
    
    # UNCOLORED clean borders
    tblBorders = parse_xml(r'''<w:tblBorders %s>
        <w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>
        <w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>
        <w:insideH w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>
        <w:insideV w:val="none"/>
        <w:left w:val="none"/>
        <w:right w:val="none"/>
    </w:tblBorders>''' % nsdecls('w'))
    tblPr.append(tblBorders)
    
    TOTAL_WIDTH = 6.11
    if not col_widths:
        max_lens = [len(h) for h in headers]
        for r in rows:
            for j, c in enumerate(r):
                max_lens[j] = max(max_lens[j], len(str(c)))
        weights = [max(l ** 0.65, 3.0) for l in max_lens]
        total_w = sum(weights)
        col_widths = [Inches((w / total_w) * TOTAL_WIDTH) for w in weights]
        
    base_sz = 8.0 if num_cols >= 7 else (8.5 if num_cols >= 5 else 9.0)
    
    tr_h = tbl.rows[0]._tr
    tr_h.get_or_add_trPr().append(parse_xml(r'<w:tblHeader %s/>' % nsdecls('w')))
    tr_h.get_or_add_trPr().append(parse_xml(r'<w:cantSplit %s/>' % nsdecls('w')))
    
    for j, h_text in enumerate(headers):
        cell = tbl.cell(0, j)
        cell.width = col_widths[j]
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p_c = cell.paragraphs[0]
        p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_c.paragraph_format.space_before = Pt(0)
        p_c.paragraph_format.space_after = Pt(0)
        p_c.paragraph_format.line_spacing = 1.05
        render_inline(p_c, h_text, base_size=base_sz)
        for r in p_c.runs:
            r.bold = True
            r.font.color.rgb = RGBColor(0, 0, 0)
            
    for j in range(num_cols):
        tcPr = tbl.cell(0, j)._tc.get_or_add_tcPr()
        tcBorders = parse_xml(r'''<w:tcBorders %s>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>
        </w:tcBorders>''' % nsdecls('w'))
        tcPr.append(tcBorders)
        
    for i, row_data in enumerate(rows):
        tr = tbl.rows[i+1]._tr
        tr.get_or_add_trPr().append(parse_xml(r'<w:cantSplit %s/>' % nsdecls('w')))
        for j, cell_text in enumerate(row_data):
            cell = tbl.cell(i+1, j)
            cell.width = col_widths[j]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p_c = cell.paragraphs[0]
            p_c.paragraph_format.space_before = Pt(0)
            p_c.paragraph_format.space_after = Pt(0)
            p_c.paragraph_format.line_spacing = 1.05
            
            is_num = bool(re.match(r'^[+\-]?[0-9.,%S= ()±]+$', str(cell_text).strip()))
            p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER if is_num else WD_ALIGN_PARAGRAPH.LEFT
            render_inline(p_c, str(cell_text), base_size=base_sz)
            
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(6)

def add_reference(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4.0)
    p.paragraph_format.left_indent = Inches(0.35)
    p.paragraph_format.first_line_indent = Inches(-0.35)
    render_inline(p, text, base_size=10.0)
